# report_docx.py — 生成与原校对报告排版一致的 Word 报告
# 排版还原：标题 Arial 20pt 加粗 深红(C00000)；小节用 Heading 2；
# 概览表标签列灰底(EFEFEF)；统计/问题表表头蓝底(D5E8F0)；问题表 5 列。
from datetime import date
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Arial"
HDR_FILL = "D5E8F0"   # 表头蓝
LBL_FILL = "EFEFEF"   # 概览标签列灰
TITLE_RED = RGBColor(0xC0, 0x00, 0x00)


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def _set(cell, text, bold=False, size=10, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    r = p.add_run(text if text is not None else "")
    r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold
    return cell


def _widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w


def _kv_table(doc, rows):
    from docx.shared import Inches
    t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
    for k, v in rows:
        c = t.add_row().cells
        _set(c[0], k, bold=True); _shade(c[0], LBL_FILL)
        _set(c[1], v)
    _widths(t, [Inches(1.9), Inches(4.6)])
    return t


def _issue_table(doc, issues):
    from docx.shared import Inches
    cols = ["#", "位置", "问题 / 原文", "类型", "修改建议 / 说明"]
    t = doc.add_table(rows=1, cols=5); t.style = "Table Grid"
    for i, h in enumerate(cols):
        _set(t.rows[0].cells[i], h, bold=True); _shade(t.rows[0].cells[i], HDR_FILL)
    for it in issues:
        loc = (it.get("lang") or "—")
        if it.get("line"):
            loc += f" · L{it['line']}"
        c = t.add_row().cells
        _set(c[0], it["code"]); _set(c[1], loc); _set(c[2], it["text"])
        _set(c[3], it["kind"]); _set(c[4], it["note"])
    _widths(t, [Inches(0.7), Inches(1.25), Inches(2.5), Inches(1.15), Inches(2.1)])
    return t


def _h2(doc, text):
    h = doc.add_heading(level=2)
    r = h.add_run(text); r.font.name = FONT
    return h


def build(meta, result, out_path):
    doc = Document()
    # 默认字体
    n = doc.styles["Normal"]; n.font.name = FONT; n.font.size = Pt(10)

    # 标题
    p = doc.add_paragraph()
    r = p.add_run("校对报告")
    r.font.name = FONT; r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = TITLE_RED
    sub = doc.add_paragraph()
    rs = sub.add_run(meta.get("product", meta.get("filename", "")))
    rs.font.name = FONT; rs.font.size = Pt(11); rs.font.bold = True

    st = result["stats"]
    # 一、报告概览
    _h2(doc, "一、报告概览")
    _kv_table(doc, [
        ("输入来源", meta.get("filename", "")),
        ("产品 / 物料类型", meta.get("material", "说明书 / 用户手册（说明书类）")),
        ("风格标准", "Apple 技术文档风格（专业、精准、简洁、无营销腔）"),
        ("OCR 语言覆盖", meta.get("langs", "")),
        ("检查方式", "OCR → 结构化 → 自动机械检查（线索层，不替代视觉终审）"),
        ("制表日期", meta.get("date", date.today().isoformat())),
    ])
    doc.add_paragraph()
    # 问题统计
    from docx.shared import Inches
    t = doc.add_table(rows=2, cols=4); t.style = "Table Grid"
    for i, h in enumerate(["问题总数", "高优先", "疑似串版（需人工确认）", "低优先（单位/风格）"]):
        _set(t.rows[0].cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER); _shade(t.rows[0].cells[i], HDR_FILL)
    for i, v in enumerate([st["total"], st["high"], st["warn"], st["low"]]):
        _set(t.rows[1].cells[i], str(v), align=WD_ALIGN_PARAGRAPH.CENTER)
    _widths(t, [Inches(1.6), Inches(1.6), Inches(2.0), Inches(1.7)])

    if result.get("audit_summary"):
        _h2(doc, "AI 审计结论（独立审计员对全面性与准确性的把关）")
        ap = doc.add_paragraph()
        ar = ap.add_run(result["audit_summary"])
        ar.font.name = FONT; ar.font.size = Pt(10); ar.font.italic = True

    high = [i for i in result["issues"] if i["severity"] == "high"]
    warn = [i for i in result["issues"] if i["severity"] == "warn"]
    low = [i for i in result["issues"] if i["severity"] == "low"]

    _h2(doc, "二、高优先问题（参数自洽 / 事实待核实）")
    if high:
        _issue_table(doc, high)
    else:
        doc.add_paragraph("（无）")

    _h2(doc, "三、疑似串版（高召回 · 需人工逐条确认）")
    if warn:
        _issue_table(doc, warn)
    else:
        doc.add_paragraph("（无）")

    _h2(doc, "四、低优先（单位空格 / 风格禁用词）")
    if low:
        _issue_table(doc, low)
    else:
        doc.add_paragraph("（无）")

    _h2(doc, "五、说明与边界")
    for line in [
        "本报告为自动机械检查的「线索层」，不替代视觉终审。",
        "遵循四条铁律：只校产品内容 / 禁止编造 / OCR 不纠错 / 截断不补全。",
        "OCR 可能错字（型号、对勾矩阵、特殊符号），数值/枚举类仅作定位，最终以矢量原稿视觉确认为准。",
        "疑似串版为高召回候选，需人工确认（ES/IT 易混对已标低置信）；语义型交叉引用（标题↔机型）与纯视觉乱码本层无法识别。",
    ]:
        b = doc.add_paragraph(style="List Bullet"); rr = b.add_run(line); rr.font.name = FONT; rr.font.size = Pt(10)

    doc.save(out_path)
    return out_path
