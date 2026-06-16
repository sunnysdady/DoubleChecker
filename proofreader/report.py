#!/usr/bin/env python3
# report.py — 由合并后的 findings 生成交付报告，套用 VCD 校对报告模板
#  ① Word(.docx)：横向 A4·Arial·红黑灰极简，黑底白字表头、粉底红字「类型」列、按语言分节
#  ② 自包含 HTML：同款排版，离线可看
# 用法： python3 report.py merged.json out_dir   （merged.json = {meta, issues, stats, audit_summary}）
import sys, json, html as _h
from datetime import date

# 模板配色
RED = "DC1E1E"; DARK = "1D1D1F"; GRAY = "6E6E73"; AMBER = "B25E00"; GREEN = "1A7F37"
HDR_BG = "1A1A1A"; BANNER_BG = "FFF7E8"; TYPE_BG = "FFF0F0"; WHITE = "FFFFFF"
COLW = [0.43, 1.04, 3.54, 1.26, 4.17]            # 5 列宽（英寸，横向 A4）
LANG_NAMES = [("EN", "英语"), ("ES", "西班牙语"), ("DE", "德语"),
              ("FR", "法语"), ("IT", "意大利语"), ("JP", "日语")]
NAME = dict(LANG_NAMES)


def _sec_of(it):
    """归属：高优先/结构性 → 结构性表；否则按语言分节。"""
    s = it.get("section")
    if s:
        return s
    if it.get("severity") == "high":
        return "结构性"
    lang = (it.get("lang") or "").strip()
    if "←" in lang:
        return lang.split("←")[0].strip()        # 目标语言版本
    return lang if lang in NAME else "其他"


def _loc(it):
    return it.get("loc") or (f"L{it['line']}" if it.get("line") else (it.get("lang") or "—"))


def _group(issues):
    structural = [i for i in issues if _sec_of(i) == "结构性"]
    bylang = {}
    for i in issues:
        s = _sec_of(i)
        if s == "结构性":
            continue
        bylang.setdefault(s, []).append(i)
    return structural, bylang


# ================================================================ Word
def build_docx(meta, result, out_path):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
        from docx.enum.section import WD_ORIENT
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        return None
    FONT = "Arial"

    def rgb(hexs):
        return RGBColor(int(hexs[0:2], 16), int(hexs[2:4], 16), int(hexs[4:6], 16))

    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

    def setc(cell, text, size=9, bold=False, color=DARK, fill=None, align=None):
        cell.text = ""; p = cell.paragraphs[0]
        if align:
            p.alignment = align
        r = p.add_run(text if text is not None else "")
        r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(color)
        if fill:
            shade(cell, fill)

    def para(text, size, bold=False, color=DARK, space_after=2):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space_after)
        r = p.add_run(text); r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = rgb(color)
        return p

    def issue_table(items):
        t = doc.add_table(rows=1, cols=5); t.style = "Table Grid"; t.autofit = False
        for i, hname in enumerate(["#", "位置", "问题", "类型", "修改建议"]):
            setc(t.rows[0].cells[i], hname, size=9, bold=True, color=WHITE, fill=HDR_BG)
        for n, it in enumerate(items, 1):
            c = t.add_row().cells
            setc(c[0], str(n), size=9, bold=True, color=GRAY, align=AL.CENTER)
            setc(c[1], _loc(it), size=8, color=GRAY)
            setc(c[2], it.get("text", ""), size=9, color=DARK)
            setc(c[3], it.get("kind", ""), size=8, bold=True, color=RED, fill=TYPE_BG)
            setc(c[4], it.get("note", ""), size=9, color=DARK)
        for row in t.rows:
            for i, w in enumerate(COLW):
                row.cells[i].width = Inches(w)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11.69), Inches(8.27)
    sec.left_margin = sec.right_margin = Inches(0.62)
    try:
        doc.styles["Normal"].font.name = FONT; doc.styles["Normal"].font.size = Pt(9)
    except KeyError:
        pass

    st = result.get("stats", {})
    langs = meta.get("langs", "")
    # 头部
    para("校对报告 · PROOFREADING REPORT", 9, bold=True, color=RED, space_after=1)
    para(meta.get("product", meta.get("filename", "")), 20, bold=True, color=DARK, space_after=2)
    sub = " 　·　 ".join(x for x in [meta.get("filename", ""), meta.get("pages", ""),
                                     "说明书类（Apple 技术文档风格）",
                                     meta.get("date", date.today().isoformat())] if x)
    para(sub, 9, color=GRAY, space_after=2)
    para(f"语言覆盖：{langs}", 9, bold=True, color=DARK, space_after=6)

    # 统计四格
    structural, bylang = _group(result.get("issues", []))
    lang_count = meta.get("lang_count") or len(bylang) or "—"
    tiles = [(st.get("total", 0), "问题总数", DARK), (st.get("high", 0), "高优先", AMBER),
             (st.get("warn", 0), "待确认 / 疑似", AMBER),
             (lang_count, "语言版本", DARK)]
    tt = doc.add_table(rows=1, cols=4); tt.style = "Table Grid"
    for i, (num, label, color) in enumerate(tiles):
        cell = tt.rows[0].cells[i]; cell.text = ""; p = cell.paragraphs[0]; p.alignment = AL.CENTER
        r1 = p.add_run(str(num)); r1.font.name = FONT; r1.font.size = Pt(18); r1.font.bold = True
        r1.font.color.rgb = rgb(color); r1.add_break()
        r2 = p.add_run(label); r2.font.name = FONT; r2.font.size = Pt(9); r2.font.color.rgb = rgb(GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 审计结论
    if result.get("audit_summary"):
        bt = doc.add_table(rows=1, cols=1); bt.style = "Table Grid"
        setc(bt.rows[0].cells[0], "✔ AI 审计结论（独立审计员把关）： " + result["audit_summary"],
             size=9, bold=True, color="5B3A8C", fill="F4F0FF")
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 结构性 / 高优先
    if structural:
        bt = doc.add_table(rows=1, cols=1); bt.style = "Table Grid"
        setc(bt.rows[0].cells[0], "⚠ 结构性 / 跨语言高优先问题（建议设计与产品侧优先处理）",
             size=9.5, bold=True, color=AMBER, fill=BANNER_BG)
        issue_table(structural)

    # 按语言分节
    for code, name in LANG_NAMES:
        items = bylang.get(code)
        if not items:
            continue
        para(f"{code}  {name}    · {len(items)} 处", 12, bold=True, color=RED, space_after=2)
        issue_table(items)
    for other, items in bylang.items():
        if other in NAME:
            continue
        para(f"{other}    · {len(items)} 处", 12, bold=True, color=RED, space_after=2)
        issue_table(items)

    # 边界
    para("说明与边界", 11, bold=True, color=DARK, space_after=2)
    for line in ["本报告为线索层，不替代视觉终审。",
                 "四条铁律：只校产品内容 / 禁止编造 / OCR 不纠错 / 截断不补全。",
                 "事实真伪、纯视觉乱码、语义型交叉引用需人工 + 联网最终确认。"]:
        p = doc.add_paragraph(line, style="List Bullet")
        for r in p.runs:
            r.font.name = FONT; r.font.size = Pt(8); r.font.color.rgb = rgb(GRAY)
    doc.save(out_path)
    return out_path


# ================================================================ HTML（同款排版）
def _rows_html(items):
    r = ['<table><tr><th>#</th><th>位置</th><th>问题</th><th>类型</th><th>修改建议</th></tr>']
    for n, it in enumerate(items, 1):
        r.append(
            f'<tr><td class="i">{n}</td><td class="loc">{_h.escape(_loc(it))}</td>'
            f'<td>{_h.escape(it.get("text",""))}</td>'
            f'<td class="ty">{_h.escape(it.get("kind",""))}</td>'
            f'<td>{_h.escape(it.get("note",""))}</td></tr>')
    r.append('</table>')
    return "".join(r)


def build_html(meta, result, out_path):
    st = result.get("stats", {})
    structural, bylang = _group(result.get("issues", []))
    body = ""
    if result.get("audit_summary"):
        body += f'<div class="audit"><b>✔ AI 审计结论（独立审计员把关）：</b>{_h.escape(result["audit_summary"])}</div>'
    if structural:
        body += '<div class="banner">⚠ 结构性 / 跨语言高优先问题（建议设计与产品侧优先处理）</div>' + _rows_html(structural)
    for code, name in LANG_NAMES:
        items = bylang.get(code)
        if items:
            body += f'<h2>{code}　{name}　· {len(items)} 处</h2>' + _rows_html(items)
    for other, items in bylang.items():
        if other not in NAME:
            body += f'<h2>{_h.escape(other)}　· {len(items)} 处</h2>' + _rows_html(items)
    tiles = [(st.get("total", 0), "问题总数", DARK), (st.get("high", 0), "高优先", AMBER),
             (st.get("warn", 0), "待确认 / 疑似", AMBER),
             (meta.get("lang_count") or len(bylang) or "—", "语言版本", DARK)]
    tilehtml = "".join(f'<div><b style="color:#{c}">{n}</b><span>{_h.escape(l)}</span></div>'
                       for n, l, c in tiles)
    sub = " 　·　 ".join(x for x in [meta.get("filename", ""), meta.get("pages", ""),
                                     "说明书类（Apple 技术文档风格）",
                                     meta.get("date", date.today().isoformat())] if x)
    doc = f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">
<title>校对报告 · {_h.escape(meta.get('product',''))}</title><style>
body{{font-family:Arial,"PingFang SC","Microsoft YaHei",sans-serif;max-width:1100px;margin:0 auto;padding:28px 22px;color:#1d1d1f;line-height:1.6}}
.kicker{{color:#{RED};font-weight:700;font-size:13px;letter-spacing:.04em}}
h1{{font-size:30px;margin:2px 0;color:#1d1d1f}}
.sub{{color:#{GRAY};font-size:13px;margin:2px 0}} .cov{{font-weight:700;font-size:13px;margin:4px 0 14px}}
.tiles{{display:flex;gap:10px;margin:12px 0 18px}} .tiles div{{border:1px solid #e5e5e7;border-radius:10px;padding:10px 22px;text-align:center}}
.tiles b{{font-size:26px;display:block;line-height:1.1}} .tiles span{{font-size:12px;color:#{GRAY}}}
.banner{{background:#{BANNER_BG};color:#{AMBER};font-weight:700;font-size:13px;padding:8px 12px;border-radius:6px;margin:16px 0 8px}}
.audit{{background:#f4f0ff;border-left:4px solid #5b3a8c;padding:10px 14px;border-radius:0 8px 8px 0;margin:12px 0;font-size:13px}}
h2{{color:#{RED};font-size:17px;margin:22px 0 6px}}
table{{width:100%;border-collapse:collapse;margin:4px 0 8px;font-size:13px}}
th{{background:#{HDR_BG};color:#fff;font-weight:700;padding:6px 8px;text-align:left;font-size:12px}}
td{{border:1px solid #e5e5e7;padding:6px 8px;vertical-align:top}}
td.i{{color:#{GRAY};font-weight:700;text-align:center;width:32px}} td.loc{{color:#{GRAY};font-size:12px;white-space:nowrap}}
td.ty{{background:#{TYPE_BG};color:#{RED};font-weight:700;font-size:12px;white-space:nowrap}}
.boundary{{margin-top:22px;color:#{GRAY};font-size:12px;border-top:1px solid #e5e5e7;padding-top:10px}}
</style></head><body>
<div class="kicker">校对报告 · PROOFREADING REPORT</div>
<h1>{_h.escape(meta.get('product', meta.get('filename','')))}</h1>
<div class="sub">{_h.escape(sub)}</div>
<div class="cov">语言覆盖：{_h.escape(meta.get('langs',''))}</div>
<div class="tiles">{tilehtml}</div>
{body}
<div class="boundary">本报告为线索层，不替代视觉终审。四条铁律：只校产品内容 / 禁止编造 / OCR 不纠错 / 截断不补全。
事实真伪、纯视觉乱码、语义型交叉引用需人工 + 联网最终确认。</div>
</body></html>"""
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(doc)
    return out_path


if __name__ == "__main__":
    data = json.load(open(sys.argv[1], encoding="utf-8"))
    out_dir = sys.argv[2] if len(sys.argv) > 2 else "."
    meta = data.get("meta", {})
    result = {"issues": data.get("issues", []), "stats": data.get("stats", {}),
              "audit_summary": data.get("audit_summary", "")}
    print("✓ HTML:", build_html(meta, result, f"{out_dir}/report.html"))
    d = build_docx(meta, result, f"{out_dir}/report.docx")
    print("✓ Word:", d) if d else print("· Word 跳过（pip install python-docx）")
